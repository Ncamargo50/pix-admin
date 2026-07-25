# -*- coding: utf-8 -*-
"""Flujo de caja Pixadvisor — modelo editable con formulas (openpyxl)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"C:/Users/Usuario/Desktop/Flujo_de_Caja_Pixadvisor.xlsx"
FN='Arial'
TEAL='0D9488'; TEALD='0F766E'; LIME='7FD633'; LIMET='EFFAD9'; TEALT='E8F6F4'
DARK='0F172A'; GRAY='64748B'; BLUE='0000FF'; GREEN='008000'; WHITE='FFFFFF'; SURF='F8FAFC'
NUM='#,##0;[Red](#,##0);"-"'
thin=Side(style='thin',color='D9D9D9')
BORD=Border(left=thin,right=thin,top=thin,bottom=thin)

def C(ws,ref,val=None,bold=False,color='000000',size=10,fill=None,align=None,fmt=None,italic=False,bd=False,wrap=False):
    c=ws[ref]
    if val is not None: c.value=val
    c.font=Font(name=FN,bold=bold,color=color,size=size,italic=italic)
    if fill: c.fill=PatternFill('solid',fgColor=fill)
    c.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    if fmt: c.number_format=fmt
    if bd: c.border=BORD
    return c

def title(ws,ref_end,text,sub):
    ws.merge_cells('A1:%s1'%ref_end); ws.merge_cells('A2:%s2'%ref_end)
    C(ws,'A1',text,bold=True,color=WHITE,size=14,fill=TEAL,align='left')
    C(ws,'A2',sub,color=WHITE,size=9,fill=TEALD,align='left')
    ws.row_dimensions[1].height=24; ws.row_dimensions[2].height=16

def section(ws,row,ref_end,text):
    ws.merge_cells('A%d:%s%d'%(row,ref_end,row))
    C(ws,'A%d'%row,text,bold=True,color=TEALD,size=10.5,fill=TEALT,align='left')
    ws.row_dimensions[row].height=18

wb=Workbook()

# ============================================================ PARAMETROS
ws=wb.active; ws.title='Parametros'
ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=40; ws.column_dimensions['B'].width=16; ws.column_dimensions['C'].width=34
title(ws,'C','PIXADVISOR  —  Parámetros del Flujo de Caja','Agricultura de Precisión · todos los valores en Bolivianos (Bs) salvo indicación')

section(ws,4,'C','TIPOS DE CAMBIO')
C(ws,'A5','Tipo de cambio USD → Bs',bd=True); C(ws,'B5',6.96,color=BLUE,align='right',fmt='#,##0.00',bd=True); C(ws,'C5','Oficial',italic=True,color=GRAY,bd=True)
C(ws,'A6','Tipo de cambio R$ → Bs',bd=True); C(ws,'B6',2.10,color=BLUE,align='right',fmt='#,##0.00',bd=True); C(ws,'C6','Acordado con Nilton',italic=True,color=GRAY,bd=True)

section(ws,8,'C','EGRESOS MENSUALES (Bs)')
rows=[('Retirada por socio (Bs/mes)',23000),('Número de socios',2),('Técnico de campo — Romelio Vaca Cuellar',4500),
      ('Viático semanal (Bs)',2500),('Semanas por mes','=52/12'),('APRISA (pago fijo)',2800),
      ('Mantenimiento camioneta',1000),('Internet / agua / luz / teléfono',1500),('Despensas extras',3000)]
r=9
for name,val in rows:
    C(ws,'A%d'%r,name,bd=True)
    if isinstance(val,str):
        C(ws,'B%d'%r,val,color='000000',align='right',fmt='#,##0.00',bd=True)   # formula 52/12
    else:
        C(ws,'B%d'%r,val,color=BLUE,align='right',fmt=('#,##0' if name!='Número de socios' else '0'),bd=True)
    r+=1
# r is now 18

section(ws,19,'C','POLÍTICA FINANCIERA')
C(ws,'A20','Meses de subsistencia (colchón)',bd=True); C(ws,'B20',5,color=BLUE,align='right',fmt='0',bd=True)
C(ws,'A21','Retiro anual extra por socio (Año 3+)',bd=True); C(ws,'B21',350000,color=BLUE,align='right',fmt=NUM,bd=True)
C(ws,'A22','Saldo inicial de caja',bd=True); C(ws,'B22',0,color=BLUE,align='right',fmt=NUM,bd=True)

section(ws,24,'C','INVERSIONES')
C(ws,'A25','Vehículos — cantidad',bd=True); C(ws,'B25',2,color=BLUE,align='right',fmt='0',bd=True)
C(ws,'A26','Vehículos — precio (USD c/u)',bd=True); C(ws,'B26',25000,color=BLUE,align='right',fmt='#,##0',bd=True)
C(ws,'A27','Drones — cantidad',bd=True); C(ws,'B27',2,color=BLUE,align='right',fmt='0',bd=True)
C(ws,'A28','Drones — precio (R$ c/u)',bd=True); C(ws,'B28',85000,color=BLUE,align='right',fmt='#,##0',bd=True)
C(ws,'A29','Casa en anticrético (USD)',bd=True); C(ws,'B29',35000,color=BLUE,align='right',fmt='#,##0',bd=True)

section(ws,31,'C','VALORES CALCULADOS (no editar)')
C(ws,'A32','Egreso mensual total (Bs)',bold=True,bd=True)
C(ws,'B32','=B9*B10+B11+B12*B13+B14+B15+B16+B17',bold=True,align='right',fmt=NUM,bd=True)
C(ws,'A33','Egreso anual total (Bs)',bold=True,bd=True); C(ws,'B33','=B32*12',bold=True,align='right',fmt=NUM,bd=True)
C(ws,'A34','Fondo de subsistencia (Bs)',bold=True,bd=True); C(ws,'B34','=B32*B20',bold=True,align='right',fmt=NUM,bd=True)

section(ws,36,'C','LEYENDA DE COLORES')
C(ws,'A37','Azul = dato editable / supuesto',color=BLUE,size=9)
C(ws,'A38','Negro = fórmula dentro de la hoja',color='000000',size=9)
C(ws,'A39','Verde = enlace a otra hoja',color=GREEN,size=9)

# helper to reference parametros
P=lambda ref:"'Parametros'!%s"%ref

# ============================================================ FLUJO MENSUAL
ws=wb.create_sheet('Flujo Mensual')
ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=32
for i in range(2,14): ws.column_dimensions[get_column_letter(i)].width=11
ws.column_dimensions['N'].width=13
title(ws,'N','PIXADVISOR  —  Flujo de Caja Mensual (Año 1)','Cargue los ingresos reales en la fila azul; los egresos se calculan solos · valores en Bs')

# header row 4
C(ws,'A4','Concepto',bold=True,color=WHITE,fill=TEALD,bd=True)
for i in range(12):
    C(ws,'%s4'%get_column_letter(2+i),'Mes %d'%(i+1),bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
C(ws,'N4','Total Año',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)

def fill_row(row,label,formula_func,color='000000',bold=False,fill=None,total='sum'):
    C(ws,'A%d'%row,label,bold=bold,fill=fill,bd=True)
    for i in range(12):
        col=get_column_letter(2+i)
        C(ws,'%s%d'%(col,row),formula_func(col),color=color,bold=bold,align='right',fmt=NUM,fill=fill,bd=True)
    if total=='sum':
        C(ws,'N%d'%row,'=SUM(B%d:M%d)'%(row,row),bold=True,align='right',fmt=NUM,fill=fill,bd=True)

# Ingresos (input, blue) row 5
C(ws,'A5','INGRESOS por servicios y ventas',bold=True,fill=SURF,bd=True)
for i in range(12):
    col=get_column_letter(2+i)
    C(ws,'%s5'%col,143000,color=BLUE,align='right',fmt=NUM,fill=SURF,bd=True)
C(ws,'N5','=SUM(B5:M5)',bold=True,align='right',fmt=NUM,fill=SURF,bd=True)

# Egresos header
section(ws,6,'N','EGRESOS OPERATIVOS')
fill_row(7,'Retiradas socios',lambda c:'=%s*%s'%(P('$B$9'),P('$B$10')),color=GREEN)
fill_row(8,'Técnico Romelio Vaca Cuellar',lambda c:'=%s'%P('$B$11'),color=GREEN)
fill_row(9,'Viáticos de campo',lambda c:'=%s*%s'%(P('$B$12'),P('$B$13')),color=GREEN)
fill_row(10,'APRISA (pago fijo)',lambda c:'=%s'%P('$B$14'),color=GREEN)
fill_row(11,'Mantenimiento camioneta',lambda c:'=%s'%P('$B$15'),color=GREEN)
fill_row(12,'Internet / agua / luz / teléfono',lambda c:'=%s'%P('$B$16'),color=GREEN)
fill_row(13,'Despensas extras',lambda c:'=%s'%P('$B$17'),color=GREEN)
fill_row(14,'Total Egresos Operativos',lambda c:'=SUM(%s7:%s13)'%(c,c),bold=True,fill=TEALT)

# Inversiones / capital (input)
C(ws,'A15','Inversiones / egresos de capital',italic=True,bd=True)
for i in range(12):
    col=get_column_letter(2+i); C(ws,'%s15'%col,0,color=BLUE,align='right',fmt=NUM,bd=True)
C(ws,'N15','=SUM(B15:M15)',bold=True,align='right',fmt=NUM,bd=True)

# Flujo neto
C(ws,'A16','Flujo neto del mes',bold=True,fill=LIMET,bd=True)
for i in range(12):
    col=get_column_letter(2+i)
    C(ws,'%s16'%col,'=%s5-%s14-%s15'%(col,col,col),bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
C(ws,'N16','=SUM(B16:M16)',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)

# Saldo acumulado
C(ws,'A17','Saldo acumulado de caja',bold=True,fill=SURF,bd=True)
C(ws,'B17','=%s+B16'%P('$B$22'),bold=True,align='right',fmt=NUM,fill=SURF,bd=True)
for i in range(1,12):
    col=get_column_letter(2+i); prev=get_column_letter(1+i)
    C(ws,'%s17'%col,'=%s17+%s16'%(prev,col),bold=True,align='right',fmt=NUM,fill=SURF,bd=True)
C(ws,'N17','=M17',bold=True,align='right',fmt=NUM,fill=SURF,bd=True)

# Fondo objetivo (referencia)
C(ws,'A18','Fondo de subsistencia objetivo',italic=True,color=GRAY,bd=True)
for i in range(12):
    col=get_column_letter(2+i)
    C(ws,'%s18'%col,'=%s'%P('$B$34'),color=GREEN,italic=True,align='right',fmt=NUM,bd=True)
C(ws,'N18','=%s'%P('$B$34'),color=GREEN,italic=True,align='right',fmt=NUM,bd=True)

C(ws,'A20','Nota: «Mes 1…12» es genérico — renómbralo a tus meses reales. Ingresos en azul = placeholder (meta Año 1); reemplázalo por lo facturado.',italic=True,color=GRAY,size=8)

# ============================================================ PROYECCION ANUAL
ws=wb.create_sheet('Proyeccion Anual')
ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=42
for c in ['B','C','D']: ws.column_dimensions[c].width=16
title(ws,'D','PIXADVISOR  —  Proyección 3 Años','Plan de inversiones e ingreso requerido por año · valores en Bs')

section(ws,4,'D','PLAN DE INVERSIONES (Bs)')
C(ws,'A5','Vehículo (1 unidad)',bd=True); C(ws,'B5','=%s*%s'%(P('$B$26'),P('$B$5')),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A6','Dron de mapeo (1 unidad)',bd=True); C(ws,'B6','=%s*%s'%(P('$B$28'),P('$B$6')),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A7','Casa en anticrético',bd=True); C(ws,'B7','=%s*%s'%(P('$B$29'),P('$B$5')),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A8','Inversión TOTAL (2 vehíc. + 2 drones + casa)',bold=True,bd=True); C(ws,'B8','=2*B5+2*B6+B7',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
C(ws,'A9','Asignación Año 1 (1 vehículo + 2 drones)',bd=True); C(ws,'B9','=B5+2*B6',align='right',fmt=NUM,bd=True)
C(ws,'A10','Asignación Año 2 (1 vehículo + casa)',bd=True); C(ws,'B10','=B5+B7',align='right',fmt=NUM,bd=True)

section(ws,12,'D','PROYECCIÓN ANUAL — INGRESO REQUERIDO')
C(ws,'A13','Concepto',bold=True,color=WHITE,fill=TEALD,bd=True)
C(ws,'B13','Año 1',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
C(ws,'C13','Año 2',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
C(ws,'D13','Año 3+',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
# rows
C(ws,'A14','Egresos operativos del año',bd=True)
for col in ['B','C','D']: C(ws,'%s14'%col,'=%s'%P('$B$33'),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A15','Constitución fondo de subsistencia',bd=True)
C(ws,'B15','=%s'%P('$B$34'),color=GREEN,align='right',fmt=NUM,bd=True); C(ws,'C15',0,align='right',fmt=NUM,bd=True); C(ws,'D15',0,align='right',fmt=NUM,bd=True)
C(ws,'A16','Inversiones del año',bd=True)
C(ws,'B16','=B9',align='right',fmt=NUM,bd=True); C(ws,'C16','=B10',align='right',fmt=NUM,bd=True); C(ws,'D16',0,align='right',fmt=NUM,bd=True)
C(ws,'A17','Retiros extraordinarios socios (2 × 350.000)',bd=True)
C(ws,'B17',0,align='right',fmt=NUM,bd=True); C(ws,'C17',0,align='right',fmt=NUM,bd=True)
C(ws,'D17','=%s*%s'%(P('$B$21'),P('$B$10')),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A18','INGRESO REQUERIDO (Bs/año)',bold=True,fill=LIMET,bd=True)
for col in ['B','C','D']: C(ws,'%s18'%col,'=SUM(%s14:%s17)'%(col,col),bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
C(ws,'A19','Ingreso requerido (Bs/mes)',bold=True,bd=True)
for col in ['B','C','D']: C(ws,'%s19'%col,'=%s18/12'%col,bold=True,align='right',fmt=NUM,bd=True)

C(ws,'A21','Nota: Año 1 es el más exigente (cubre operación + fondo + 1ª etapa de inversión). Los retiros de Bs 350.000/socio arrancan en el Año 3, ya con las inversiones hechas.',italic=True,color=GRAY,size=8,wrap=True)
ws.merge_cells('A21:D22')

# ============================================================ ANALISIS
ws=wb.create_sheet('Analisis')
ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=46; ws.column_dimensions['B'].width=16; ws.column_dimensions['C'].width=16
title(ws,'C','PIXADVISOR  —  Análisis y Metas','Fondo de subsistencia · punto de equilibrio · metas de facturación')

section(ws,4,'C','FONDO DE SUBSISTENCIA (colchón)')
C(ws,'A5','Egreso mensual total',bd=True); C(ws,'B5','=%s'%P('$B$32'),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A6','Meses objetivo',bd=True); C(ws,'B6','=%s'%P('$B$20'),color=GREEN,align='right',fmt='0',bd=True)
C(ws,'A7','FONDO NECESARIO (Bs)',bold=True,bd=True); C(ws,'B7','=B5*B6',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
C(ws,'A8','Alternativa: solo operativo (sin retiradas socios)',bd=True); C(ws,'B8','=(%s-%s*%s)*%s'%(P('$B$32'),P('$B$9'),P('$B$10'),P('$B$20')),align='right',fmt=NUM,bd=True)

section(ws,10,'C','PUNTO DE EQUILIBRIO')
C(ws,'A11','Facturación mínima mensual (cubre costos)',bd=True); C(ws,'B11','=%s'%P('$B$32'),color=GREEN,align='right',fmt=NUM,bd=True)
C(ws,'A12','Facturación mínima anual (cubre costos)',bd=True); C(ws,'B12','=%s'%P('$B$33'),color=GREEN,align='right',fmt=NUM,bd=True)

section(ws,14,'C','METAS DE FACTURACIÓN POR ESCENARIO')
C(ws,'A15','Escenario',bold=True,color=WHITE,fill=TEALD,bd=True)
C(ws,'B15','Bs / año',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
C(ws,'C15','Bs / mes',bold=True,color=WHITE,fill=TEALD,align='center',bd=True)
metas=[('1 · Equilibrio operativo (solo cubrir costos)',"='Proyeccion Anual'!$B$14"),
       ('2 · Equilibrio + fondo de subsistencia',"='Proyeccion Anual'!$B$14+'Proyeccion Anual'!$B$15"),
       ('3 · Año 1 completo (operación + fondo + inversión)',"='Proyeccion Anual'!$B$18"),
       ('4 · Año 2 (operación + inversión)',"='Proyeccion Anual'!$C$18"),
       ('5 · Régimen objetivo Año 3+ (con retiros socios)',"='Proyeccion Anual'!$D$18")]
r=16
for name,f in metas:
    C(ws,'A%d'%r,name,bd=True)
    C(ws,'B%d'%r,f,color=GREEN,align='right',fmt=NUM,bd=True)
    C(ws,'C%d'%r,'=B%d/12'%r,align='right',fmt=NUM,bd=True)
    r+=1
# emphasize regime row
for col in ['A','B','C']: ws['%s20'%col].fill=PatternFill('solid',fgColor=LIMET)
ws['A20'].font=Font(name=FN,bold=True); ws['B20'].font=Font(name=FN,bold=True,color=GREEN); ws['C20'].font=Font(name=FN,bold=True)

section(ws,22,'C','RECOMENDACIONES')
recs=['Prioridad 1: constituir el fondo de subsistencia de Bs 348.167 (5 meses) ANTES de invertir — protege ante paros, falta de combustible o meses flojos.',
      'El Año 1 es el más exigente: apuntar a ~Bs 143.000/mes de facturación para cubrir operación, fondo y la primera etapa de inversión.',
      'Comprar primero los 2 drones (herramienta que genera ingresos) y 1 vehículo; la casa en anticrético y el 2º vehículo en el Año 2.',
      'Desde el Año 3, con inversiones hechas, el negocio debe facturar ~Bs 1.535.600/año para sostener operación + retiros de Bs 350.000 por socio.',
      'Revisar el tipo de cambio USD: si no consiguen dólares al oficial (6,96), las inversiones suben fuerte — actualizar la celda en Parámetros.']
r=23
for t in recs:
    ws.merge_cells('A%d:C%d'%(r,r))
    C(ws,'A%d'%r,'•  '+t,size=9,wrap=True,align='left')
    ws.row_dimensions[r].height=28
    r+=1

# freeze panes
wb['Flujo Mensual'].freeze_panes='B5'
wb['Parametros'].freeze_panes='A3'

# forzar recalculo al abrir (Excel/Calc) ya que no hay LibreOffice para precalcular aqui
wb.calculation.fullCalcOnLoad = True

wb.save(OUT)
print('saved', OUT)
