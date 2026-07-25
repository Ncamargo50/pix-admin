# -*- coding: utf-8 -*-
"""Agrega hoja 'Metas por Producto' (4 productos, incl. On-Farm) al flujo de caja."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

PATH=r"C:/Users/Usuario/Desktop/PIXADVISOR_Flujo_de_Caja_2026-05-31/Flujo_de_Caja_Pixadvisor.xlsx"
FN='Arial'
TEAL='0D9488'; TEALD='0F766E'; LIMET='EFFAD9'; TEALT='E8F6F4'
GRAY='64748B'; BLUE='0000FF'; GREEN='008000'; WHITE='FFFFFF'
NUM='#,##0;[Red](#,##0);"-"'; USD='"$"#,##0.00'; PCT='0%'; DEC='#,##0.00'; UNI='#,##0'
thin=Side(style='thin',color='D9D9D9'); BORD=Border(left=thin,right=thin,top=thin,bottom=thin)

wb=load_workbook(PATH)
if 'Metas por Producto' in wb.sheetnames: del wb['Metas por Producto']
ws=wb.create_sheet('Metas por Producto')
ws.sheet_view.showGridLines=False
ws.column_dimensions['A'].width=36
for c in ['B','C','D','E','F']: ws.column_dimensions[c].width=15

def C(ref,val=None,bold=False,color='000000',size=10,fill=None,align=None,fmt=None,italic=False,bd=False,wrap=False):
    c=ws[ref]
    if val is not None: c.value=val
    c.font=Font(name=FN,bold=bold,color=color,size=size,italic=italic)
    if fill: c.fill=PatternFill('solid',fgColor=fill)
    c.alignment=Alignment(horizontal=align,vertical='center',wrap_text=wrap)
    if fmt: c.number_format=fmt
    if bd: c.border=BORD
    return c
def sec(row,text,end='F'):
    ws.merge_cells('A%d:%s%d'%(row,end,row)); C('A%d'%row,text,bold=True,color=TEALD,size=10.5,fill=TEALT,align='left'); ws.row_dimensions[row].height=18
def hdr(row,labels):
    for i,l in enumerate(labels):
        C('%s%d'%(chr(65+i),row),l,bold=True,color=WHITE,fill=TEALD,align=('left' if i==0 else 'right'),bd=True)

ws.merge_cells('A1:F1'); ws.merge_cells('A2:F2')
C('A1','PIXADVISOR  —  Metas de Venta por Producto',bold=True,color=WHITE,size=14,fill=TEAL,align='left')
C('A2','Cuánto vender de cada producto para llegar a la meta · 4 productos (incl. On-Farm)',color=WHITE,size=9,fill=TEALD,align='left')
ws.row_dimensions[1].height=24; ws.row_dimensions[2].height=16

sec(4,'PARÁMETROS DE VENTA')
C('A5','Cambio paralelo USD → Bs',bd=True); C('B5',10,color=BLUE,align='right',fmt=DEC,bd=True)
ws.merge_cells('C5:F5'); C('C5','(On-Farm ya está en Bs; el cambio aplica a MAX PIROL, análisis y mapeo)',italic=True,color=GRAY,size=9)

sec(7,'PRODUCTOS Y GANANCIA UNITARIA')
hdr(8,['Producto','Precio USD','Margen','Precio Bs','Ganancia Bs/u'])
# (name, usd_price o None, margen, precio_bs_directo o None)
prods=[('MAX PIROL (por litro)',8,0.25,None),
       ('Análisis de suelo (por ha)',15,0.20,None),
       ('Mapeamento (por ha)',10,0.35,None),
       ('On-Farm biofertilizante (por litro)',None,0.15,79.79)]
for i,(name,usd,mg,bsp) in enumerate(prods):
    r=9+i; C('A%d'%r,name,bd=True)
    if usd is None: C('B%d'%r,'—',align='right',color=GRAY,bd=True)
    else: C('B%d'%r,usd,color=BLUE,align='right',fmt=USD,bd=True)
    C('C%d'%r,mg,color=BLUE,align='right',fmt=PCT,bd=True)
    if bsp is None: C('D%d'%r,'=B%d*$B$5'%r,align='right',fmt=DEC,bd=True)
    else: C('D%d'%r,bsp,color=BLUE,align='right',fmt=DEC,bd=True)
    C('E%d'%r,'=D%d*C%d'%(r,r),bold=True,align='right',fmt=DEC,bd=True)
# rows 9,10,11,12

sec(14,'META DE MARGEN ANUAL (Bs)')
C('A15','Equilibrio — cubrir todos los costos',bd=True); C('B15',"='Proyeccion Anual'!$B$14",color=GREEN,align='right',fmt=NUM,bd=True)
C('A16','Meta plena Año 3+ (con retiros socios)',bd=True); C('B16',"='Proyeccion Anual'!$D$18",color=GREEN,align='right',fmt=NUM,bd=True)
C('A17','META USADA EN ESTE ANÁLISIS',bold=True,bd=True); C('B17','=B16',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
ws.merge_cells('C17:F17'); C('C17','← cambiá a =B15 para analizar el equilibrio',italic=True,color=GRAY,size=9)

sec(19,'OPCIÓN A — cada producto SOLO para llegar a la meta')
hdr(20,['Producto','Ganancia Bs/u','Unid./año','Unid./mes','Facturación Bs/año'])
amap=[('MAX PIROL (litros)',9),('Análisis de suelo (ha)',10),('Mapeamento (ha)',11),('On-Farm (litros)',12)]
for i,(name,pr) in enumerate(amap):
    r=21+i; C('A%d'%r,name,bd=True)
    C('B%d'%r,'=E%d'%pr,color=GREEN,align='right',fmt=DEC,bd=True)
    C('C%d'%r,'=$B$17/B%d'%r,bold=True,align='right',fmt=UNI,bd=True)
    C('D%d'%r,'=C%d/12'%r,align='right',fmt=UNI,bd=True)
    C('E%d'%r,'=C%d*D%d'%(r,pr),align='right',fmt=NUM,bd=True)
# rows 21,22,23,24

sec(26,'OPCIÓN B — mezcla recomendada (más peso a mapeo y MAX PIROL)')
hdr(27,['Producto','% aporte','Margen aportado Bs','Unid./año','Unid./mes','Facturación Bs/año'])
bmap=[('MAX PIROL (litros)',9,0.30),('Análisis de suelo (ha)',10,0.20),('Mapeamento (ha)',11,0.40),('On-Farm (litros)',12,0.10)]
for i,(name,pr,share) in enumerate(bmap):
    r=28+i; C('A%d'%r,name,bd=True)
    C('B%d'%r,share,color=BLUE,align='right',fmt=PCT,bd=True)
    C('C%d'%r,'=B%d*$B$17'%r,align='right',fmt=NUM,bd=True)
    C('D%d'%r,'=C%d/E%d'%(r,pr),bold=True,align='right',fmt=UNI,bd=True)
    C('E%d'%r,'=D%d/12'%r,align='right',fmt=UNI,bd=True)
    C('F%d'%r,'=D%d*D%d'%(r,pr),align='right',fmt=NUM,bd=True)
# rows 28,29,30,31
C('A32','TOTAL',bold=True,fill=LIMET,bd=True)
C('B32','=SUM(B28:B31)',bold=True,align='right',fmt=PCT,fill=LIMET,bd=True)
C('C32','=SUM(C28:C31)',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
C('D32','',fill=LIMET,bd=True); C('E32','',fill=LIMET,bd=True)
C('F32','=SUM(F28:F31)',bold=True,align='right',fmt=NUM,fill=LIMET,bd=True)
ws.merge_cells('A33:F33'); C('A33','Mezcla recomendada: mapeo 40% · MAX PIROL 30% · análisis 20% · On-Farm 10% (suma 100%). Cambiá los % (azul) para simular tu mezcla. On-Farm: Bs 79,79/L, 15% (trato Gil).',italic=True,color=GRAY,size=8,wrap=True)
ws.row_dimensions[33].height=24

wb.calculation.fullCalcOnLoad=True
wb.save(PATH)
print('OK 4 productos. Hojas:', wb.sheetnames)

def bs(x): return ('Bs {:,.0f}'.format(x)).replace(',','.')
def n(x): return '{:,.0f}'.format(x).replace(',','.')
rate=10; meta=1535600
items={'MAX PIROL':(8*rate,0.25),'Analisis':(15*rate,0.20),'Mapeo':(10*rate,0.35),'On-Farm':(79.79,0.15)}
print('--- Ganancia Bs/u ---')
for k,(p,m) in items.items(): print('   %-10s precio %s  gana %s'%(k,bs(p),'Bs %.2f'%(p*m)))
print('--- OPCION A (meta plena %s) ---'%bs(meta))
for k,(p,m) in items.items():
    g=p*m; u=meta/g; print('   %-10s %s u/año (%s u/mes) | factura %s'%(k,n(u),n(u/12),bs(u*p)))
shares={'MAX PIROL':0.30,'Analisis':0.20,'Mapeo':0.40,'On-Farm':0.10}
print('--- OPCION B (mezcla recomendada 40/30/20/10) ---'); tot=0
for k,(p,m) in items.items():
    g=p*m; mm=meta*shares[k]; u=mm/g; f=u*p; tot+=f
    print('   %-10s aporte %s | %s u/año (%s u/mes) | factura %s'%(k,bs(mm),n(u),n(u/12),bs(f)))
print('   TOTAL facturacion:', bs(tot))
